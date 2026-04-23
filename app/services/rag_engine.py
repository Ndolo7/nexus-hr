import io
import logging
import re
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
from sqlalchemy import create_engine, or_
from app.models.db_models import HRDocument
from app.core.config import settings

try:
    from langchain_google_genai import GoogleGenerativeAIEmbeddings
    has_langchain = True
except ImportError:
    has_langchain = False

logger = logging.getLogger(__name__)

# Initialize engine globally for the service
engine = create_engine(settings.DATABASE_URL)


class RAGEngine:
    def __init__(self):
        self.logger = logging.getLogger(__name__)
        self._embedding_model = settings.GOOGLE_EMBEDDING_MODEL
        if has_langchain:
            self.embeddings = self._build_embeddings(self._embedding_model)
        else:
            self.embeddings = None
            self.logger.warning("langchain_google_genai not installed. Cannot perform real embedding generation.")

    def _build_embeddings(self, model_name: str):
        return GoogleGenerativeAIEmbeddings(
            model=model_name,
            google_api_key=settings.GOOGLE_API_KEY,
        )

    def _embedding_model_candidates(self) -> List[str]:
        configured = (settings.GOOGLE_EMBEDDING_MODEL or "").strip()
        configured_base = configured.replace("models/", "") if configured else ""
        candidates = [
            configured,
            f"models/{configured_base}" if configured_base else "",
            configured_base,
            "models/text-embedding-004",
            "text-embedding-004",
            "models/embedding-001",
            "embedding-001",
        ]
        # Preserve order while de-duplicating.
        return list(dict.fromkeys([c for c in candidates if c]))

    def _embed_query(self, text: str) -> List[float]:
        if not self.embeddings:
            raise RuntimeError("Embeddings are not configured.")

        last_error: Optional[Exception] = None
        for candidate in self._embedding_model_candidates():
            try:
                if candidate != self._embedding_model:
                    self.logger.warning(
                        "Switching embedding model from '%s' to '%s'.",
                        self._embedding_model,
                        candidate,
                    )
                    self.embeddings = self._build_embeddings(candidate)
                    self._embedding_model = candidate

                return self.embeddings.embed_query(text)
            except Exception as e:
                last_error = e
                error_text = str(e).lower()
                model_missing = (
                    "not found" in error_text
                    or "not supported for embedcontent" in error_text
                    or "'code': 404" in error_text
                    or "status': 'not_found'" in error_text
                )
                if model_missing:
                    self.logger.warning(
                        "Embedding model '%s' unavailable. Trying another candidate.",
                        candidate,
                    )
                    continue
                raise

        raise RuntimeError(
            f"Unable to generate embeddings with available models: {self._embedding_model_candidates()}"
        ) from last_error

    def _score_keyword_match(self, query: str, title: str, content: str) -> float:
        """Lightweight lexical scoring fallback when vector search is unavailable."""
        tokens = [
            t for t in re.findall(r"[a-z0-9]+", query.lower())
            if len(t) > 2 and t not in {"the", "and", "for", "with", "from", "that", "this", "policy"}
        ]
        if not tokens:
            return 0.0

        title_lower = (title or "").lower()
        content_lower = (content or "").lower()
        score = 0.0

        for token in tokens:
            score += 3.0 * title_lower.count(token)
            score += 1.0 * content_lower.count(token)

        return score

    def _keyword_fallback_search(self, session: Session, query: str, top_k: int) -> List[Dict[str, Any]]:
        """Fallback policy search using SQL ILIKE + simple keyword scoring."""
        terms = [t for t in re.findall(r"[a-z0-9]+", query.lower()) if len(t) > 2]
        conditions = []
        for term in terms:
            like_expr = f"%{term}%"
            conditions.append(HRDocument.title.ilike(like_expr))
            conditions.append(HRDocument.content.ilike(like_expr))

        candidate_query = session.query(HRDocument)
        if conditions:
            candidate_query = candidate_query.filter(or_(*conditions))

        candidates = candidate_query.limit(max(top_k * 8, 20)).all()
        scored = [
            {
                "title": doc.title,
                "content": doc.content,
                "relevance_score": self._score_keyword_match(query, doc.title, doc.content),
            }
            for doc in candidates
        ]
        scored.sort(key=lambda d: d["relevance_score"], reverse=True)
        return scored[:top_k]

    # ------------------------------------------------------------------
    # Text extraction
    # ------------------------------------------------------------------
    def extract_text(self, raw_bytes: bytes, filename: str) -> str:
        """Extract plain text from a TXT, PDF, or DOCX file."""
        filename_lower = filename.lower()

        if filename_lower.endswith(".txt"):
            return raw_bytes.decode("utf-8", errors="replace")

        elif filename_lower.endswith(".pdf"):
            try:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(raw_bytes))
                return "\n".join(page.extract_text() or "" for page in reader.pages)
            except ImportError:
                raise RuntimeError("pypdf is not installed. Run: pip install pypdf")

        elif filename_lower.endswith(".docx"):
            try:
                from docx import Document
                doc = Document(io.BytesIO(raw_bytes))
                return "\n".join(para.text for para in doc.paragraphs)
            except ImportError:
                raise RuntimeError("python-docx is not installed. Run: pip install python-docx")

        else:
            raise ValueError(f"Unsupported file extension: {filename}")

    # ------------------------------------------------------------------
    # Ingest
    # ------------------------------------------------------------------
    def ingest_document(self, title: str, content: str, category: str) -> int:
        """Vectorise content and persist it to the hr_documents table."""
        embedding_vector: Optional[List[float]] = None

        if self.embeddings:
            try:
                embedding_vector = self._embed_query(content)
            except Exception as e:
                self.logger.warning(f"Embedding generation failed: {e}. Storing without vector.")

        with Session(engine) as session:
            doc = HRDocument(
                title=title,
                content=content,
                category=category,
                embedding=embedding_vector,
            )
            session.add(doc)
            session.commit()
            session.refresh(doc)
            self.logger.info(f"Indexed document '{title}' with id={doc.id}")
            return doc.id

    # ------------------------------------------------------------------
    # Retrieval
    # ------------------------------------------------------------------
    def retrieve_policy(self, query: str, top_k: int = 3) -> List[Dict[str, Any]]:
        """
        Retrieve policies with vector search first, then keyword fallback.
        """
        self.logger.info("Retrieving policies for: %s", query)

        with Session(engine) as session:
            if self.embeddings:
                try:
                    query_embedding = self._embed_query(query)
                    results = (
                        session.query(HRDocument)
                        .filter(HRDocument.embedding.isnot(None))
                        .order_by(HRDocument.embedding.cosine_distance(query_embedding))
                        .limit(top_k)
                        .all()
                    )
                    if results:
                        return [{"title": r.title, "content": r.content} for r in results]
                except Exception as e:
                    session.rollback()
                    self.logger.warning(
                        "Vector policy retrieval failed. Falling back to keyword search. Error: %s",
                        e,
                    )
            else:
                self.logger.info("Embeddings unavailable. Using keyword fallback search.")

            return self._keyword_fallback_search(session, query, top_k)

    def format_retrieved_context(self, docs: List[Dict[str, Any]]) -> str:
        """Format the retrieved documents to inject into LLM prompt."""
        if not docs:
            return "No relevant company policy documents found."
        context = ""
        for i, doc in enumerate(docs):
            context += f"Document {i+1}: {doc['title']}\nContent: {doc['content']}\n\n"
        return context


rag_engine = RAGEngine()
